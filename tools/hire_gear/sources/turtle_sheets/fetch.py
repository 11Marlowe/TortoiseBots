#!/usr/bin/env python3
"""Fetch Turtle class workbooks (xlsx export, hyperlinks = Turtle item IDs).

Books (from raw research): 9 per-class workbooks; each gear tab is named
"Naxx <spec> BIS" / "K40 <spec> BIS". Gear cells hyperlink to
database.turtlecraft.gg/?item=<id> = exact Turtle item IDs.

Writes: ../../data/turtle_sheets.json
Row: {class, tab, list_name, slot, item_name, item_id, enchant}
"""
import json
import re
import sys
import urllib.request
from pathlib import Path

SRC = Path(__file__).resolve().parent
OUT = SRC / ".." / ".." / "data" / "turtle_sheets.json"
UA = {"User-Agent": "TortoiseBots-hire-gear-research/1.0 (offline BIS snapshot)"}

BOOKS = {
    "druid": "1KfDPyRLOeBEymO9poDDdOEeFgjig_cMc0L1TxaiJOKI",
    "hunter": "1uwIMjtOPwNNYS9hxekeqT22zRktOVIwMoa1NnKXnG8o",
    "mage": "1zFqF3BRNSsFUjiRY5PW4bmsHKSSFPxzRqudi61oOzeM",
    "paladin": "1CNna2XjnJ5WbNDfHVL94E8uQvIKZwfJhuEC7HTJE76A",
    "priest": "1v-8RQKGYFOVvb9Ew8ZMayoC6eX56aY2H10Xq3-l8AJw",
    "rogue": "1ueQMvqDxWXf1TiByARJAPNg0f5-XbIhcLDBPDoM0-Ts",
    "shaman": "18pUBSdqpeWKz9PBkj8_n-bReYS3aUs9AXCRSztQHitI",
    "warlock": "1hiykADgWDZI-baXI_-AVZYD8dwwR31w-x5jUnXp2u30",
    "warrior": "1ji0xqex39AibBsSI_Wpaov39yDTZeJFaAMJ88vstDY8",
}
COVER = "1gBYi4RsHW6Ddjx06qQEUZPWmXAQRX3UvBH0uwLOugdA"

LINK_RE = re.compile(r"[?&]item=(\d+)")
GEAR_SLOTS = {"Head", "Neck", "Shoulders", "Back", "Chest", "Wrist", "Hands",
              "Waist", "Legs", "Feet", "Ring 1", "Ring 2", "Trinket 1",
              "Trinket 2", "Main Hand", "Offhand", "Ranged", "Relic",
              "Trinket Swap"}


def fetch_xlsx(sid, dest):
    req = urllib.request.Request(
        f"https://docs.google.com/spreadsheets/d/{sid}/export?format=xlsx",
        headers=UA)
    with urllib.request.urlopen(req, timeout=300) as r, open(dest, "wb") as f:
        f.write(r.read())


def link_id(target):
    if not target:
        return None
    m = LINK_RE.search(target)
    return int(m.group(1)) if m else None


def main():
    import openpyxl
    only = sys.argv[1:]
    tabs_meta, rows = {}, []
    cache = SRC / ".xlsx-cache"
    cache.mkdir(exist_ok=True)
    for cls, sid in BOOKS.items():
        if only and cls not in only:
            continue
        xlsx = cache / f"{cls}.xlsx"
        if not xlsx.exists():
            print(f"{cls}: downloading xlsx…", flush=True)
            fetch_xlsx(sid, xlsx)
        wb = openpyxl.load_workbook(xlsx, data_only=True)
        tabs_meta[cls] = list(wb.sheetnames)
        for tab in wb.sheetnames:
            if "BIS" not in tab:
                continue
            ws = wb[tab]
            section, ench, gear, swaps = "gear", {}, [], []
            for row in ws.iter_rows(min_col=2, max_col=3):
                lab = row[0].value
                cell = row[1]
                name = cell.value
                if lab is None and name is None:
                    continue
                lab = str(lab).strip() if lab is not None else ""
                if lab in ("Trinket Swaps",):
                    section = "swaps"
                    continue
                if lab in ("Enchants",):
                    section = "enchants"
                    continue
                if lab in ("Set Bonus", "Talents", "Gear Stats", "Strength",
                           "Consumables List", "Raid Buffs", ""):
                    continue
                if not lab or not name:
                    continue
                name = str(name).strip()
                iid = link_id(cell.hyperlink.target if cell.hyperlink else None)
                if section == "enchants":
                    ench[lab] = name
                elif section == "swaps" or lab == "Trinket Swap":
                    swaps.append((name, iid))
                elif lab in GEAR_SLOTS:
                    gear.append((lab, name, iid))
            lst = "K40" if tab.startswith("K40") else "Naxx"
            for lab, name, iid in gear:
                rows.append({"class": cls, "tab": tab, "list_name": lst,
                             "slot": lab, "item_name": name, "item_id": iid,
                             "enchant": ench.get(lab)})
            for name, iid in swaps:
                rows.append({"class": cls, "tab": tab, "list_name": lst,
                             "slot": "Trinket Swap", "item_name": name,
                             "item_id": iid, "enchant": None})
            print(f"{cls}/{tab}: {len(gear)} gear {len(swaps)} swaps"
                  f" {len(ench)} enchants", flush=True)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps({"source": "turtle-sheets",
                               "cover": COVER, "books": BOOKS,
                               "fetched": "2026-10-12",
                               "tabs": tabs_meta, "rows": rows}, indent=1) + "\n")
    noid = sum(1 for r in rows if r["item_id"] is None)
    print(f"rows={len(rows)} missing_id={noid} -> {OUT}")


if __name__ == "__main__":
    sys.exit(main())
